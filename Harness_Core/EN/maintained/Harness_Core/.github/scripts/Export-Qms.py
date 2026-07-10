#!/usr/bin/env python3
"""
Export-Qms.py
=============
Convert QMS Markdown documents to Word (.docx) using PDLM reference templates.

Workflow per document:
  1. Validate  — no unresolved placeholders; has RECORD CHANGE SUMMARY;
                 passes Verify-QmsStructure.py (skeleton conformance)
  2. Mermaid   — pre-render every ```mermaid block to PNG via mmdc
  3. Pandoc    — convert processed markdown to .docx (heading levels shifted -1
                 so ## becomes Heading 1, ### becomes Heading 2, etc.)
  4. Cover     — prepend the template cover page to the pandoc-generated body

Usage:
    python .github/scripts/Export-Qms.py [options] <file.md> [<file.md> ...]

Options:
    --output-dir DIR        Output directory            (default: each source .md file's folder)
    --templates-dir DIR     PDLM templates directory    (default: docs/qms-templates)
    --mermaid-timeout SEC   Per-diagram timeout seconds (default: 600)
"""

import argparse
import copy
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph as DocxParagraph
except ImportError:
    print(
        "ERROR: python-docx is not installed.\n"
        "       Run:  pip install python-docx",
        file=sys.stderr,
    )
    sys.exit(1)


# ---------------------------------------------------------------------------
# Template map — markdown basename → PDLM reference template filename
# ---------------------------------------------------------------------------
TEMPLATE_MAP: dict[str, str] = {
    "SwRS.md":        "2003000201 HA030-054-05_PDLM SwRS Template.docx",
    "SSDS.md":        "2003000198 HA030-054-02_PDLM SDS_SSDS Template.docx",
    "SDD.md":         "2003000204 HA030-055-02_PDLM Software Design Document Template _.docx",
    "MVP.md":         "2003000256 HA030-064-01_PDLM Module Verification or Module Integration Plan Template.docx",
    "MVProcedure.md": "2003000258 HA030-064-03_PDLM Module Verification or Module Integration Procedure Template.docx",
    "MVReport.md":    "2003000257 HA030-064-02_PDLM Module Verification or Module Integration Report Template.docx",
    # --- System-level (black-box) verification, owned by @test-designer. ---
    # PDLM .docx templates are not in the repo yet; the filenames below are the
    # expected drop-in names.  Until the .docx are added these stems are skipped
    # gracefully (see PENDING_STEMS).  Replace the filenames if the official
    # template names differ once the templates are added.
    "VerificationPlan.md":      "PDLM Software Verification Plan Template.docx",
    "VerificationProcedure.md": "PDLM Software Verification Procedure Template.docx",
}

# Stems whose PDLM .docx template has not been added yet: a missing template is
# a graceful SKIP, not a hard error.  Remove a stem once its .docx is in place.
PENDING_STEMS: set[str] = {"VerificationPlan", "VerificationProcedure"}


# ---------------------------------------------------------------------------
# Tool discovery
# ---------------------------------------------------------------------------
def _find_exe(name: str, hint: str) -> Path:
    """
    Locate an executable on PATH.  On Windows, .cmd wrappers (npm CLIs) take
    priority over .ps1 — they can be invoked directly via cmd.exe /c.
    """
    for candidate in (name + ".cmd", name + ".exe", name):
        found = shutil.which(candidate)
        if found:
            return Path(found)
    raise RuntimeError(f"'{name}' not found on PATH.  {hint}")


def check_prerequisites() -> tuple[Path, Path]:
    pandoc = _find_exe("pandoc", "Install: winget install pandoc")
    mmdc   = _find_exe("mmdc",   "Install: npm install -g @mermaid-js/mermaid-cli")
    return pandoc, mmdc


# ---------------------------------------------------------------------------
# Subprocess runner
# ---------------------------------------------------------------------------
def _run(cmd: list, timeout: int, cwd: str | None = None) -> tuple[str, str]:
    """
    Execute a command list.  .cmd / .bat files are wrapped in cmd.exe /c so
    Windows resolves the extension correctly without shell=True.
    Raises RuntimeError on non-zero exit or timeout.
    """
    str_cmd = [str(a) for a in cmd]
    exe = str_cmd[0].lower()
    if exe.endswith((".cmd", ".bat")):
        str_cmd = ["cmd.exe", "/c"] + str_cmd

    result = subprocess.run(
        str_cmd, capture_output=True, text=True, timeout=timeout, cwd=cwd
    )
    if result.returncode != 0:
        raise RuntimeError(
            f"Command failed (exit {result.returncode}):\n"
            f"  {' '.join(str_cmd)}\n"
            f"stderr: {result.stderr.strip()}\n"
            f"stdout: {result.stdout.strip()}"
        )
    return result.stdout, result.stderr


# ---------------------------------------------------------------------------
# Document validation
# ---------------------------------------------------------------------------
_RE_FENCED  = re.compile(r"```.*?```", re.DOTALL)
_RE_INLINE  = re.compile(r"`[^`\r\n]+`")
_RE_COMMENT = re.compile(r"<!--.*?-->", re.DOTALL)
_RE_HOLDER  = re.compile(r"\{[A-Za-z0-9_ -]+\}|<[^>]*template[^>]*>")


def validate_document(path: Path, content: str) -> None:
    """Fail fast if the markdown has unresolved placeholders or no change record."""
    prose = _RE_FENCED.sub("", content)
    prose = _RE_INLINE.sub("", prose)
    prose = _RE_COMMENT.sub("", prose)

    hits = _RE_HOLDER.findall(prose)
    if hits:
        sample = ", ".join(hits[:5])
        raise ValueError(f"{path.name}: unresolved placeholders — {sample}")

    if "RECORD CHANGE SUMMARY" not in content:
        raise ValueError(f"{path.name}: missing RECORD CHANGE SUMMARY section.")


def verify_structure(md_path: Path, tmpl_dir: Path) -> None:
    """
    Enforce skeleton conformance before converting.  Delegates to the sibling
    Verify-QmsStructure.py (the single source of truth for the rules) so the
    exported .docx can never drift from the PDLM template structure.
    Raises RuntimeError if the document fails verification.
    """
    verifier = Path(__file__).with_name("Verify-QmsStructure.py")
    if not verifier.exists():
        raise RuntimeError(f"structure verifier not found: {verifier}")

    cmd = [
        sys.executable, str(verifier),
        "--check-content",
        "--living-dir", str(md_path.parent),
        "--skeleton-dir", str(tmpl_dir / "skeletons"),
        md_path.name,
    ]
    proc = subprocess.run(cmd, capture_output=True, text=True)
    if proc.returncode != 0:
        detail = (proc.stdout or "") + (proc.stderr or "")
        raise RuntimeError(
            f"{md_path.name}: structure verification failed —\n{detail.strip()}"
        )



# ---------------------------------------------------------------------------
# Mermaid pre-rendering
# ---------------------------------------------------------------------------
_RE_MERMAID = re.compile(r"```mermaid\r?\n(.*?)\r?\n```", re.DOTALL)


def render_mermaid(
    content: str,
    doc_name: str,
    work: Path,
    puppeteer_cfg: Path,
    mmdc: Path,
    timeout: int,
) -> str:
    """Replace every ```mermaid block with a PNG image reference."""
    matches = list(_RE_MERMAID.finditer(content))
    if not matches:
        return content

    print(f"  {doc_name}: rendering {len(matches)} Mermaid diagram(s)…")

    parts: list[str] = []
    cursor = 0
    for i, m in enumerate(matches, start=1):
        parts.append(content[cursor : m.start()])

        mmd_path = work / f"{doc_name}-{i}.mmd"
        png_path = work / f"{doc_name}-{i}.png"
        mmd_path.write_text(m.group(1), encoding="utf-8")

        try:
            _run(
                [
                    mmdc,
                    "--input",              mmd_path,
                    "--output",             png_path,
                    "--backgroundColor",    "white",
                    "--scale",              "2",
                    "--puppeteerConfigFile", puppeteer_cfg,
                ],
                timeout=timeout,
                cwd=str(work),
            )
        except Exception as exc:
            raise RuntimeError(f"Mermaid diagram #{i} in {doc_name} failed:\n{exc}") from exc

        if not png_path.exists():
            raise RuntimeError(f"Mermaid diagram #{i} in {doc_name}: no PNG produced.")

        parts.append(f"![{doc_name} diagram {i}]({png_path})")
        cursor = m.end()

    parts.append(content[cursor:])
    return "".join(parts)


# ---------------------------------------------------------------------------
# Cover-page merge
# ---------------------------------------------------------------------------
def _get_heading_numid(template: Document) -> str | None:
    """
    Return the numId that drives multi-level heading numbering in the template.

    Resolution order:
    1. First Heading 1 paragraph with an explicit <w:numPr> override in the
       template body (e.g. numId=26 in the PDLM SwRS template).
    2. The Heading 1 *style* numId, but only if that abstractNum has a
       <w:pStyle w:val="Heading1"/> link on ilvl=0 — confirming it is the
       intended heading outline list (e.g. numId=26 in the PDLM SSDS template).

    Returns None if neither heuristic succeeds.
    """
    # 1. Explicit paragraph override on a Heading 1 paragraph.
    for para in template.paragraphs:
        if not para.style.name.startswith("Heading 1"):
            continue
        numPr = para._element.find(".//" + qn("w:numPr"))
        if numPr is not None:
            numId_elem = numPr.find(qn("w:numId"))
            if numId_elem is not None:
                return numId_elem.get(qn("w:val"))

    # 2. Heading 1 style numId, validated by a pStyle link in its abstractNum.
    try:
        h1_style = template.styles["Heading 1"]
        style_pPr = h1_style.element.find(".//" + qn("w:pPr"))
        if style_pPr is not None:
            style_numPr = style_pPr.find(qn("w:numPr"))
            if style_numPr is not None:
                numId_elem = style_numPr.find(qn("w:numId"))
                if numId_elem is not None:
                    numId_candidate = numId_elem.get(qn("w:val"))
                    # Validate: find abstractNum and check for pStyle=Heading1 at ilvl=0
                    root = template.part.numbering_part._element
                    for num in root.findall(qn("w:num")):
                        if num.get(qn("w:numId")) != numId_candidate:
                            continue
                        absRef = num.find(qn("w:abstractNumId"))
                        if absRef is None:
                            break
                        absId = absRef.get(qn("w:val"))
                        for absNum in root.findall(qn("w:abstractNum")):
                            if absNum.get(qn("w:abstractNumId")) != absId:
                                continue
                            for lvl in absNum.findall(qn("w:lvl")):
                                if lvl.get(qn("w:ilvl")) != "0":
                                    continue
                                ps = lvl.find(qn("w:pStyle"))
                                if ps is not None and ps.get(qn("w:val")) == "Heading1":
                                    return numId_candidate
                        break
    except (KeyError, AttributeError):
        pass

    return None


def _make_numPr(numId_val: str, ilvl_val: str):
    """Build a <w:numPr><w:ilvl .../><w:numId .../></w:numPr> element."""
    from lxml import etree
    NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    numPr = etree.SubElement(etree.Element("dummy"), f"{{{NS}}}numPr")
    ilvl  = etree.SubElement(numPr, f"{{{NS}}}ilvl")
    ilvl.set(f"{{{NS}}}val", ilvl_val)
    nid   = etree.SubElement(numPr, f"{{{NS}}}numId")
    nid.set(f"{{{NS}}}val", numId_val)
    return numPr


def apply_heading_numbering(generated: Document, template: Document) -> None:
    """
    Force all heading paragraphs in the generated document to use a single,
    unified multi-level list so that Word resolves '%1.%2.' correctly.

    Problem: Different PDLM templates wire Heading 1 and Heading 2 to *different*
    numIds (e.g. SSDS: Heading 1 → numId=26, Heading 2 → numId=19).  Since
    '%1' in a level-1 format string refers to ilvl=0 of the SAME abstractNum,
    Heading 2's %1 can never see the Heading 1 counter when they are on
    different lists.

    Fix (applied to the generated doc, not the template):
    1. Resolve the "master numId" — the single numId whose abstractNum contains
       the full multi-level outline list (identified by a pStyle=Heading1 link
       on ilvl=0, or by an explicit para override on Heading 1 in the template).
    2. In that abstractNum, ensure every ilvl 0..8 has a <w:pStyle> link to the
       corresponding heading style (Heading1..9).  This lets Word track counters
       via style membership even without paragraph-level numPr.
    3. Update every Heading-N *style* in the generated doc to use master numId
       at ilvl=N-1.  This is the critical fix for templates like SSDS where
       Heading 2 is on a different list from Heading 1.
    4. Add explicit para-level numPr(master numId, ilvl=0) to all Heading 1
       paragraphs that lack one, matching the template pattern where para
       overrides were used.
    """
    from lxml import etree
    NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    numId_val = _get_heading_numid(template)
    if not numId_val:
        return  # Cannot determine the master numId — skip.

    heading_re = re.compile(r"^Heading (\d+)$")
    num_root = generated.part.numbering_part._element

    # Resolve abstractNumId for master numId in the generated doc.
    abstract_id = None
    for num in num_root.findall(qn("w:num")):
        if num.get(qn("w:numId")) == numId_val:
            ref = num.find(qn("w:abstractNumId"))
            if ref is not None:
                abstract_id = ref.get(qn("w:val"))
            break

    # Step 2: ensure pStyle links exist for every heading level in the abstractNum.
    if abstract_id is not None:
        for abs_num in num_root.findall(qn("w:abstractNum")):
            if abs_num.get(qn("w:abstractNumId")) != abstract_id:
                continue
            for lvl in abs_num.findall(qn("w:lvl")):
                ilvl_int = int(lvl.get(qn("w:ilvl"), "0"))
                target_style = f"Heading{ilvl_int + 1}"
                existing_ps = lvl.find(qn("w:pStyle"))
                if existing_ps is None:
                    ps = etree.Element(f"{{{NS}}}pStyle")
                    ps.set(qn("w:val"), target_style)
                    # Schema order inside w:lvl: start, numFmt, pStyle, ...
                    numFmt = lvl.find(qn("w:numFmt"))
                    if numFmt is not None:
                        numFmt.addnext(ps)
                    else:
                        lvl.insert(0, ps)
            break

    # Step 3: redirect every Heading-N style to master numId at ilvl=N-1.
    # This is the primary fix for templates like SSDS where Heading 2 uses a
    # different numId from Heading 1.
    for style in generated.styles:
        m = heading_re.match(style.name)
        if not m:
            continue
        level = int(m.group(1))
        ilvl_val = str(level - 1)

        style_pPr = style.element.find(".//" + qn("w:pPr"))
        if style_pPr is None:
            style_pPr = etree.SubElement(style.element, f"{{{NS}}}pPr")

        existing_numPr = style_pPr.find(qn("w:numPr"))
        if existing_numPr is not None:
            numId_elem = existing_numPr.find(qn("w:numId"))
            if numId_elem is not None:
                numId_elem.set(qn("w:val"), numId_val)
            ilvl_elem = existing_numPr.find(qn("w:ilvl"))
            if ilvl_elem is not None:
                ilvl_elem.set(qn("w:val"), ilvl_val)
            else:
                ilvl_elem = etree.Element(f"{{{NS}}}ilvl")
                ilvl_elem.set(qn("w:val"), ilvl_val)
                existing_numPr.insert(0, ilvl_elem)
        else:
            style_pPr.append(_make_numPr(numId_val, ilvl_val))

    # Step 4: add paragraph-level numPr to Heading 1 paragraphs that lack one.
    for para in generated.paragraphs:
        if para.style.name != "Heading 1":
            continue
        if para._element.find(".//" + qn("w:numPr")) is not None:
            continue

        pPr = para._element.find(qn("w:pPr"))
        if pPr is None:
            pPr = etree.SubElement(para._element, f"{{{NS}}}pPr")
            para._element.insert(0, pPr)
        pPr.append(_make_numPr(numId_val, "0"))


def _is_heading_elem(child, template_doc: Document) -> bool:
    """Return True if *child* is a <w:p> with a Heading style."""
    if child.tag != qn("w:p"):
        return False
    para = DocxParagraph(child, template_doc)
    return (para.style.name or "").lower().startswith("heading")


def _collect_cover_elements(template_doc: Document) -> list:
    """
    Return body elements that constitute the cover page — everything before
    (and not including) the first Heading paragraph.  Tables and other block
    elements that appear before the first heading are included.
    """
    body = template_doc.element.body
    cover: list = []
    for child in body:
        if child.tag == qn("w:sectPr"):
            continue  # page-layout section properties — skip
        if _is_heading_elem(child, template_doc):
            break     # first heading signals end of cover page
        cover.append(child)
    return cover


def merge_cover_page(generated: Path, template: Path, output: Path) -> None:
    """Prepend the template cover page to the pandoc-generated document body."""
    gen_doc  = Document(str(generated))
    tmpl_doc = Document(str(template))

    # Fix outline numbering before inserting cover (so cover-page headings, if
    # any, retain their original numPr and generated headings get the correct one).
    apply_heading_numbering(gen_doc, tmpl_doc)

    cover = _collect_cover_elements(tmpl_doc)
    if not cover:
        print("  [cover] Template has no cover page; using pandoc output as-is.")
        gen_doc.save(str(output))
        return

    print(f"  [cover] Prepending {len(cover)} cover element(s) from template.")
    first = gen_doc.element.body[0]
    for elem in cover:
        first.addprevious(copy.deepcopy(elem))

    gen_doc.save(str(output))


# ---------------------------------------------------------------------------
# Single document conversion
# ---------------------------------------------------------------------------
def convert_document(
    md_path: Path,
    out_dir: Path,
    tmpl_dir: Path,
    work: Path,
    pandoc: Path,
    mmdc: Path,
    puppeteer_cfg: Path,
    mermaid_timeout: int,
) -> Path:
    name = md_path.name
    if name not in TEMPLATE_MAP:
        raise ValueError(
            f"No PDLM template mapped for '{name}'.\n"
            f"  Known files: {', '.join(sorted(TEMPLATE_MAP))}"
        )

    tmpl_path = tmpl_dir / TEMPLATE_MAP[name]
    if not tmpl_path.exists():
        raise FileNotFoundError(f"Reference template not found: {tmpl_path}")

    print(f"Converting {name} …")

    content = md_path.read_text(encoding="utf-8")
    validate_document(md_path, content)
    verify_structure(md_path, tmpl_dir)

    processed = render_mermaid(
        content, md_path.stem, work, puppeteer_cfg, mmdc, mermaid_timeout
    )

    proc_md = work / f"processed-{name}"
    proc_md.write_text(processed, encoding="utf-8")

    # Pandoc step — heading levels shifted so:
    #   # Title  → level 0  (document title / metadata, not a numbered section)
    #   ## H2    → Heading 1  (first numbered level)
    #   ### H3   → Heading 2  etc.
    body_docx = work / f"body-{md_path.stem}.docx"
    _run(
        [
            pandoc,
            proc_md,
            "--from",                    "gfm",
            "--shift-heading-level-by",  "-1",
            "--reference-doc",           tmpl_path,
            "--resource-path",           work,
            "--output",                  body_docx,
        ],
        timeout=mermaid_timeout,
    )

    final = out_dir / f"{md_path.stem}.docx"
    merge_cover_page(body_docx, tmpl_path, final)
    return final


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
def main() -> None:
    parser = argparse.ArgumentParser(
        description="Convert QMS Markdown files to Word .docx using PDLM templates."
    )
    parser.add_argument("files", nargs="+", help="Markdown files to convert")
    parser.add_argument("--output-dir",      default=None,               help="Output directory (default: same folder as each source .md file)")
    parser.add_argument("--templates-dir",   default="docs/qms-templates", help="PDLM templates directory")
    parser.add_argument("--mermaid-timeout", type=int, default=600,       help="Per-diagram timeout (s)")
    args = parser.parse_args()

    pandoc, mmdc = check_prerequisites()

    tmpl_dir = Path(args.templates_dir)

    puppeteer_cfg_content = json.dumps(
        {"args": ["--no-sandbox", "--disable-setuid-sandbox"]}
    )

    converted: list[Path] = []
    skipped: list[str] = []
    with tempfile.TemporaryDirectory(prefix="qms-export-") as tmp:
        work = Path(tmp)
        puppeteer_cfg = work / "puppeteer-config.json"
        puppeteer_cfg.write_text(puppeteer_cfg_content, encoding="utf-8")

        for file_str in args.files:
            md_path = Path(file_str)
            if not md_path.exists():
                print(f"ERROR: file not found: {md_path}", file=sys.stderr)
                sys.exit(1)

            # Pending stems are skipped gracefully while their PDLM .docx
            # template is absent — a missing template is not a hard error.
            if md_path.stem in PENDING_STEMS:
                tmpl_name = TEMPLATE_MAP.get(md_path.name)
                if tmpl_name is None or not (tmpl_dir / tmpl_name).exists():
                    print(
                        f"SKIP: {md_path.name} — PDLM template not yet available "
                        f"(pending); add '{tmpl_name}' to {tmpl_dir} to enable export."
                    )
                    skipped.append(md_path.name)
                    continue

            out_dir = Path(args.output_dir) if args.output_dir else md_path.parent
            out_dir.mkdir(parents=True, exist_ok=True)
            try:
                out = convert_document(
                    md_path, out_dir, tmpl_dir, work,
                    pandoc, mmdc, puppeteer_cfg,
                    args.mermaid_timeout,
                )
                converted.append(out)
            except Exception as exc:
                print(f"\nERROR converting {md_path.name}:\n{exc}", file=sys.stderr)
                sys.exit(1)

    print(f"\nConverted {len(converted)} document(s):")
    for p in converted:
        print(f"  {p}")
    if skipped:
        print(f"\nSkipped {len(skipped)} pending document(s) (no template yet):")
        for s in skipped:
            print(f"  {s}")


if __name__ == "__main__":
    main()
